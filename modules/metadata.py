"""Extrahiert typabhaengige Metadaten aus Bild-, Audio-, Video- und Dokumentdateien."""

import os
import re
import mimetypes


def _try(fn):
    try:
        return fn()
    except Exception:
        return None


def _clean_key_name(raw_key: str) -> str:
    """Entfernt technische Praefixe (Hachoir:, EXIF:, PDF:, DOCX:, etc.)
    und formatiert den Schluessel in ein sauberes, lesbares Format."""
    # Praefixe entfernen
    k = raw_key
    prefixes = [
        "Hachoir:", "EXIF:", "ExifRead:", "Image:", "PDF:", "DOCX:",
        "PPTX:", "XLSX:", "Tag:", "Video:", "Audio:", "Mutagen:", "Exif:"
    ]
    for p in prefixes:
        if k.startswith(p):
            k = k[len(p):].strip()
            break

    # Bekannte Exif- und PDF-Schluessel bereinigen (z.B. Image Model -> Model, /Title -> Title)
    k = k.lstrip("/")
    if k.startswith("Image "):
        k = k[6:]
    elif k.startswith("EXIF "):
        k = k[5:]

    # Snake_case in Title Case umwandeln (z.B. last_modified_by -> Last Modified By)
    if "_" in k:
        k = " ".join(part.capitalize() for part in k.split("_"))

    # CamelCase schonend trennen wenn noetig (z.B. DateTimeOriginal -> Date Time Original)
    if re.search(r"[a-z][A-Z]", k) and not any(c in k for c in [" ", "(", ")", "-", ":"]):
        k = re.sub(r"([a-z])([A-Z])", r"\1 \2", k)

    # Erstes Zeichen gross
    k = k.strip()
    if k:
        k = k[0].upper() + k[1:]
    return k


def _image_meta(path: str) -> dict:
    meta = {}
    try:
        from PIL import Image
        from PIL.ExifTags import TAGS
        import pillow_heif
        pillow_heif.register_heif_opener()
        
        img = Image.open(path)
        meta["Format"] = img.format
        meta["Größe (px)"] = f"{img.width} x {img.height}"
        meta["Farbraum"] = img.mode
        exif_raw = img._getexif()
        if exif_raw:
            for tag_id, val in exif_raw.items():
                tag = TAGS.get(tag_id, tag_id)
                meta[str(tag)] = str(val)
    except Exception:
        pass
    return meta


def _exifread_meta(path: str) -> dict:
    meta = {}
    try:
        import exifread
        with open(path, "rb") as f:
            tags = exifread.process_file(f, details=True)
        for k, v in tags.items():
            meta[str(k)] = str(v)
    except Exception:
        pass
    return meta


def _audio_meta(path: str) -> dict:
    meta = {}
    try:
        from mutagen import File as MutFile
        mf = MutFile(path)
        if mf:
            meta["Dauer (s)"] = str(round(getattr(mf.info, "length", 0), 2))
            meta["Bitrate"] = str(getattr(mf.info, "bitrate", "N/A"))
            meta["Sample Rate"] = str(getattr(mf.info, "sample_rate", "N/A"))
            tag_items = mf.tags.items() if mf.tags else []
            for k, v in tag_items:
                meta[str(k)] = str(v)
    except Exception:
        pass
    return meta


def _video_meta(path: str) -> dict:
    meta = {}
    try:
        import cv2
        cap = cv2.VideoCapture(path)
        if cap.isOpened():
            w = cap.get(cv2.CAP_PROP_FRAME_WIDTH)
            h = cap.get(cv2.CAP_PROP_FRAME_HEIGHT)
            fps = cap.get(cv2.CAP_PROP_FPS)
            frames = cap.get(cv2.CAP_PROP_FRAME_COUNT)
            
            if w > 0 and h > 0:
                meta["Auflösung"] = f"{int(w)} x {int(h)}"
            if fps > 0:
                meta["FPS"] = f"{round(fps, 2)}"
            if frames > 0 and fps > 0:
                duration = frames / fps
                meta["Videodauer (s)"] = f"{round(duration, 2)}"
                meta["Frames"] = f"{int(frames)}"
            cap.release()
    except Exception:
        pass
    return meta


def _pdf_meta(path: str) -> dict:
    meta = {}
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        info = reader.metadata
        if info:
            for k, v in info.items():
                clean_k = k.lstrip("/")
                meta[clean_k] = str(v)
        meta["Seitenanzahl"] = str(len(reader.pages))
    except Exception:
        pass
    return meta


def _docx_meta(path: str) -> dict:
    meta = {}
    try:
        from docx import Document
        doc = Document(path)
        cp = doc.core_properties
        fields = [
            "author", "category", "comments", "content_status",
            "created", "identifier", "keywords", "language",
            "last_modified_by", "last_printed", "modified",
            "revision", "subject", "title", "version"
        ]
        for f in fields:
            val = getattr(cp, f, None)
            if val is not None:
                meta[f] = str(val)
    except Exception:
        pass
    return meta


def _pptx_meta(path: str) -> dict:
    meta = {}
    try:
        from pptx import Presentation
        prs = Presentation(path)
        cp = prs.core_properties
        for attr in ["author", "created", "modified", "title", "subject", "keywords", "revision"]:
            val = getattr(cp, attr, None)
            if val is not None:
                meta[attr] = str(val)
        meta["Folienanzahl"] = str(len(prs.slides))
    except Exception:
        pass
    return meta


def _xlsx_meta(path: str) -> dict:
    meta = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        cp = wb.properties
        for attr in ["creator", "created", "modified", "lastModifiedBy", "title", "subject", "keywords", "description"]:
            val = getattr(cp, attr, None)
            if val is not None:
                meta[attr] = str(val)
        meta["Tabellenblätter"] = str(len(wb.sheetnames))
        wb.close()
    except Exception:
        pass
    return meta


def _hachoir_meta(path: str) -> dict:
    meta = {}
    try:
        from hachoir.parser import createParser
        from hachoir.metadata import extractMetadata
        parser = createParser(path)
        if parser:
            with parser:
                hm = extractMetadata(parser)
            if hm:
                for item in hm.exportPlaintext():
                    line = item.lstrip("- ")
                    if ": " in line:
                        k, v = line.split(": ", 1)
                        meta[k.strip()] = v.strip()
    except Exception:
        pass
    return meta


def get_all(path: str) -> dict:
    """Metadaten aller passenden Parser ermitteln und sauber formatiert zurueckgeben."""
    mime, _ = mimetypes.guess_type(path)
    mime = mime or ""
    raw_meta = {}

    raw_meta["MIME-Typ"] = mime or "unbekannt"

    if mime.startswith("image"):
        raw_meta.update(_image_meta(path))
        raw_meta.update(_exifread_meta(path))

    if mime.startswith("audio"):
        raw_meta.update(_audio_meta(path))
        
    if mime.startswith("video"):
        raw_meta.update(_audio_meta(path))
        raw_meta.update(_video_meta(path))

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        raw_meta.update(_pdf_meta(path))
    elif ext in (".docx",):
        raw_meta.update(_docx_meta(path))
    elif ext in (".pptx",):
        raw_meta.update(_pptx_meta(path))
    elif ext in (".xlsx",):
        raw_meta.update(_xlsx_meta(path))

    # Universellen Parser anwenden
    raw_meta.update(_hachoir_meta(path))

    # Schluessel bereinigen und leere Werte filtern
    cleaned_meta = {}
    for k, v in raw_meta.items():
        if v not in (None, "", "N/A", "None"):
            clean_k = _clean_key_name(k)
            # Nur ueberschreiben wenn noch nicht vorhanden oder vorheriger Wert leerer
            if clean_k not in cleaned_meta or (len(str(v)) > len(str(cleaned_meta[clean_k])) and cleaned_meta[clean_k] == "unbekannt"):
                cleaned_meta[clean_k] = v

    return cleaned_meta if cleaned_meta else {"Info": "Keine Metadaten gefunden"}
